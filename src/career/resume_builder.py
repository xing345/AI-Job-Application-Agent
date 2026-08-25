"""
针对性简历生成器 - 根据目标岗位生成定制简历, 并渲染为 Word (.docx)
核心价值: 简历不依赖"原简历", 而是针对选中的岗位方向/JD 量身定制
"""

import json
from pathlib import Path
from typing import List, Optional, Dict, Any
from loguru import logger

from src.utils.llm_client import LLMClient
from .schemas import UserProfile, CareerDirection, ResumeDraft, ResumeBlock

RESUME_SYSTEM_PROMPT = """你是资深的简历撰写专家。你的任务是为求职者撰写一份针对目标岗位高度定制的简历草稿。

关键原则:
1. 一切围绕目标岗位的 JD 定制: 求职意向、个人简介、技能排序、经历要点全部向目标岗位靠拢。
2. 只使用求职者提供的真实材料, 绝不编造经历、技能、业绩。
3. 把求职者已有的经历/项目重新组织、突出与目标岗位匹配的部分, 弱化无关部分。
4. 技能按与目标岗位的相关度排序。
5. 语言精炼有力, 多用动词和量化结果 (量化来自真实材料, 没有就写具体描述)。"""


class TargetedResumeGenerator:
    """针对性简历生成器"""

    def __init__(self, llm_client: Optional[LLMClient] = None):
        # 简历内容较长, 用更大 max_tokens 的客户端
        self.llm_client = llm_client or LLMClient(max_tokens=8000)

    async def generate_draft(
        self,
        profile: UserProfile,
        persona: Dict[str, Any],
        direction: CareerDirection,
        target_jobs: List[Dict],
    ) -> ResumeDraft:
        """根据目标岗位生成简历草稿"""
        # 汇总目标岗位信息
        jd_lines = []
        for i, job in enumerate(target_jobs[:5], 1):
            title = job.get("title") or job.get("url", "")
            desc = (job.get("description") or "")[:1500]
            jd_lines.append(f"[岗位{i}] {title}\n{desc}")
        jd_text = "\n\n".join(jd_lines) or "(目标岗位信息缺失, 按方向定制)"

        prompt = f"""请为求职者撰写针对目标岗位定制的简历草稿。

【求职者实际情况】
{json.dumps(profile.model_dump(), ensure_ascii=False, indent=2)}

【已生成的画像信息】
{json.dumps(persona, ensure_ascii=False, indent=2)[:3000]}

【目标岗位方向】
方向名: {direction.title}
目标职位: {direction.target_positions}
搜索关键词: {direction.keywords}
技能亮点: {direction.skill_highlights}

【目标岗位 JD】
{jd_text}

【输出要求】
生成完整 ResumeDraft JSON:
- objective: 一句话求职意向, 明确对准目标职位
- summary: 个人简介, 突出与目标岗位最匹配的优势, 2-3 句
- skills: 按相关度排序, 保留真实技能
- work_experience / projects: 每段 title/subtitle/period + bullets, bullets 突出与目标岗位匹配的职责与成果, 不编造
- education: 学校/专业/学历/时间
- certificates / languages: 真实内容, 没有就空数组
- contact_line: "邮箱 | 电话" 格式"""

        try:
            draft = await self.llm_client.generate_structured_response(
                prompt,
                ResumeDraft,
                system_prompt=RESUME_SYSTEM_PROMPT,
            )
            logger.info(f"简历草稿生成完成: {draft.name}")
            return draft
        except Exception as e:
            logger.error(f"简历草稿生成失败: {e}")
            raise

    def render_docx(self, draft: ResumeDraft, output_path: str) -> str:
        """渲染简历草稿为 Word 文档"""
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()
        self._set_document_default_font(doc)

        # ---- 姓名 + 联系方式 (居中) ----
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_run(name_p, draft.name, size=20, bold=True)

        if draft.contact_line:
            contact_p = doc.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._add_run(contact_p, draft.contact_line, size=10, color=(0x66, 0x66, 0x66))

        # ---- 求职意向 ----
        self._add_section_heading(doc, "求职意向")
        p = doc.add_paragraph()
        self._add_run(p, draft.objective, size=10.5)

        # ---- 个人简介 ----
        self._add_section_heading(doc, "个人简介")
        p = doc.add_paragraph()
        self._add_run(p, draft.summary, size=10.5)

        # ---- 核心技能 ----
        self._add_section_heading(doc, "核心技能")
        p = doc.add_paragraph()
        self._add_run(p, " · ".join(draft.skills), size=10.5)

        # ---- 工作经历 ----
        if draft.work_experience:
            self._add_section_heading(doc, "工作经历")
            for block in draft.work_experience:
                self._add_block(doc, block)

        # ---- 项目经验 ----
        if draft.projects:
            self._add_section_heading(doc, "项目经验")
            for block in draft.projects:
                self._add_block(doc, block)

        # ---- 教育背景 ----
        if draft.education:
            self._add_section_heading(doc, "教育背景")
            for block in draft.education:
                self._add_block(doc, block)

        # ---- 证书与语言 ----
        extras = []
        if draft.certificates:
            extras.append("证书: " + "、".join(draft.certificates))
        if draft.languages:
            extras.append("语言: " + "、".join(draft.languages))
        if extras:
            self._add_section_heading(doc, "其他")
            p = doc.add_paragraph()
            self._add_run(p, "\n".join(extras), size=10.5)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"简历已保存: {output_path}")
        return str(output_path)

    # ------------------------------------------------------------------ #
    def _set_document_default_font(self, doc) -> None:
        """设置文档默认中文字体"""
        from docx.oxml.ns import qn
        from docx.shared import Pt
        style = doc.styles["Normal"]
        style.font.name = "微软雅黑"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(10.5)

    def _add_run(self, paragraph, text: str, size=10.5, bold=False, color=None):
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        run = paragraph.add_run(text)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)
        return run

    def _add_section_heading(self, doc, text: str) -> None:
        """小节标题, 加粗 + 下划线分隔"""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.shared import Pt
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(12)
        run.font.bold = True
        # 段落底部加边框线
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "888888")
        pBdr.append(bottom)
        pPr.append(pBdr)
        # 段前距
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)

    def _add_block(self, doc, block: ResumeBlock) -> None:
        """一段经历/项目/教育条目"""
        from docx.shared import Pt

        # 标题行: title | subtitle    (period)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        self._add_run(p, block.title, size=11, bold=True)
        if block.subtitle:
            self._add_run(p, f"  |  {block.subtitle}", size=10.5)
        if block.period:
            period_run = p.add_run(f"    ({block.period})")
            from docx.oxml.ns import qn
            period_run.font.name = "微软雅黑"
            period_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            period_run.font.size = Pt(10)
            from docx.shared import RGBColor
            period_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # 要点列表
        for bullet in block.bullets:
            bp = doc.add_paragraph()
            bp.paragraph_format.left_indent = Pt(12)
            bp.paragraph_format.space_after = Pt(1)
            self._add_run(bp, f"• {bullet}", size=10.5)
